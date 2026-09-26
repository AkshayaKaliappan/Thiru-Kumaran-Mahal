"""
Reports app views - report calculation and data retrieval.
Report logic is separated from presentation in clean architecture.
"""
import io
from datetime import date, timedelta

from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from django.utils import timezone
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from bookings.models import Booking
from bookings.serializers import BookingSerializer
from expenses.models import Expense


def get_report_summary(queryset):
    """Calculate aggregate financial summary for a queryset of bookings."""
    agg = queryset.aggregate(
        booking_count=Count('id'),
        total_payment=Sum('total_payment'),
        received_payment=Sum('received_payment'),
        balance=Sum('balance'),
    )
    return {
        'booking_count': agg['booking_count'] or 0,
        'total_payment': float(agg['total_payment'] or 0),
        'received_payment': float(agg['received_payment'] or 0),
        'outstanding_balance': float(agg['balance'] or 0),
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def daily_report(request):
    """Daily report for a specific date."""
    report_date_str = request.query_params.get('date')
    if not report_date_str:
        report_date = timezone.localdate()
    else:
        try:
            report_date = date.fromisoformat(report_date_str)
        except ValueError:
            return Response({'error': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    # Bookings covering this date
    bookings = Booking.objects.filter(
        start_date__lte=report_date,
        end_date__gte=report_date,
    ).exclude(booking_status=Booking.STATUS_CANCELLED)

    summary = get_report_summary(bookings)
    summary['date'] = report_date.isoformat()
    summary['report_type'] = 'daily'

    # Also include expense for this date
    day_expenses = float(
        Expense.objects.filter(expense_date=report_date).aggregate(t=Sum('amount'))['t'] or 0
    )
    summary['expenses'] = day_expenses

    serializer = BookingSerializer(bookings, many=True)
    return Response({
        'summary': summary,
        'bookings': serializer.data,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def monthly_report(request):
    """Monthly report for a specific year/month."""
    year_str = request.query_params.get('year')
    month_str = request.query_params.get('month')
    today = timezone.localdate()

    try:
        year = int(year_str) if year_str else today.year
        month = int(month_str) if month_str else today.month
        if not (1 <= month <= 12):
            raise ValueError("Month must be between 1 and 12.")
    except (ValueError, TypeError) as e:
        return Response({'error': str(e)}, status=400)

    first_day = date(year, month, 1)
    if month == 12:
        last_day = date(year + 1, 1, 1) - timedelta(days=1)
    else:
        last_day = date(year, month + 1, 1) - timedelta(days=1)

    # Bookings that overlap with this month
    bookings = Booking.objects.filter(
        start_date__lte=last_day,
        end_date__gte=first_day,
    ).exclude(booking_status=Booking.STATUS_CANCELLED)

    summary = get_report_summary(bookings)
    summary['year'] = year
    summary['month'] = month
    summary['month_name'] = first_day.strftime('%B %Y')
    summary['report_type'] = 'monthly'

    # Monthly expenses
    month_expenses = float(
        Expense.objects.filter(
            expense_date__gte=first_day, expense_date__lte=last_day
        ).aggregate(t=Sum('amount'))['t'] or 0
    )
    summary['expenses'] = month_expenses
    summary['net_amount'] = summary['received_payment'] - month_expenses

    serializer = BookingSerializer(bookings, many=True)
    return Response({
        'summary': summary,
        'bookings': serializer.data,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def cancelled_report(request):
    """Report of all cancelled bookings with optional date filtering."""
    date_from_str = request.query_params.get('date_from')
    date_to_str = request.query_params.get('date_to')

    bookings = Booking.objects.filter(booking_status=Booking.STATUS_CANCELLED)

    if date_from_str:
        try:
            bookings = bookings.filter(start_date__gte=date.fromisoformat(date_from_str))
        except ValueError:
            return Response({'error': 'Invalid date_from format.'}, status=400)

    if date_to_str:
        try:
            bookings = bookings.filter(start_date__lte=date.fromisoformat(date_to_str))
        except ValueError:
            return Response({'error': 'Invalid date_to format.'}, status=400)

    summary = {
        'booking_count': bookings.count(),
        'report_type': 'cancelled',
    }

    serializer = BookingSerializer(bookings.order_by('-cancelled_at'), many=True)
    return Response({'summary': summary, 'bookings': serializer.data})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def pending_payment_report(request):
    """Report of bookings with pending/partial payments."""
    date_from_str = request.query_params.get('date_from')
    date_to_str = request.query_params.get('date_to')

    bookings = Booking.objects.filter(
        payment_status__in=[Booking.PAYMENT_NOT_PAID, Booking.PAYMENT_PARTIAL]
    ).exclude(booking_status=Booking.STATUS_CANCELLED)

    if date_from_str:
        try:
            bookings = bookings.filter(start_date__gte=date.fromisoformat(date_from_str))
        except ValueError:
            return Response({'error': 'Invalid date_from format.'}, status=400)

    if date_to_str:
        try:
            bookings = bookings.filter(start_date__lte=date.fromisoformat(date_to_str))
        except ValueError:
            return Response({'error': 'Invalid date_to format.'}, status=400)

    summary = get_report_summary(bookings)
    summary['report_type'] = 'pending_payment'

    serializer = BookingSerializer(bookings.order_by('start_date'), many=True)
    return Response({'summary': summary, 'bookings': serializer.data})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def fully_paid_report(request):
    """Report of fully paid bookings."""
    date_from_str = request.query_params.get('date_from')
    date_to_str = request.query_params.get('date_to')

    bookings = Booking.objects.filter(
        payment_status=Booking.PAYMENT_FULL
    ).exclude(booking_status=Booking.STATUS_CANCELLED)

    if date_from_str:
        try:
            bookings = bookings.filter(start_date__gte=date.fromisoformat(date_from_str))
        except ValueError:
            return Response({'error': 'Invalid date_from format.'}, status=400)

    if date_to_str:
        try:
            bookings = bookings.filter(start_date__lte=date.fromisoformat(date_to_str))
        except ValueError:
            return Response({'error': 'Invalid date_to format.'}, status=400)

    summary = get_report_summary(bookings)
    summary['report_type'] = 'fully_paid'

    serializer = BookingSerializer(bookings.order_by('-start_date'), many=True)
    return Response({'summary': summary, 'bookings': serializer.data})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_excel(request):
    """Export bookings report as Excel file."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    report_type = request.query_params.get('report_type', 'all')
    date_from_str = request.query_params.get('date_from')
    date_to_str = request.query_params.get('date_to')
    year_str = request.query_params.get('year')
    month_str = request.query_params.get('month')
    booking_status = request.query_params.get('booking_status')

    bookings = Booking.objects.all()

    # Apply filters based on report type
    if report_type == 'today':
        today = timezone.localdate()
        bookings = bookings.filter(start_date__lte=today, end_date__gte=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'upcoming':
        today = timezone.localdate()
        bookings = bookings.filter(start_date__gt=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'completed':
        today = timezone.localdate()
        bookings = bookings.filter(end_date__lt=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'cancelled':
        bookings = bookings.filter(booking_status=Booking.STATUS_CANCELLED)
    elif report_type == 'pending_payment':
        bookings = bookings.filter(
            payment_status__in=[Booking.PAYMENT_NOT_PAID, Booking.PAYMENT_PARTIAL]
        ).exclude(booking_status=Booking.STATUS_CANCELLED)
    elif report_type == 'fully_paid':
        bookings = bookings.filter(payment_status=Booking.PAYMENT_FULL).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'monthly' and year_str and month_str:
        try:
            year = int(year_str)
            month = int(month_str)
            first_day = date(year, month, 1)
            if month == 12:
                last_day = date(year + 1, 1, 1) - timedelta(days=1)
            else:
                last_day = date(year, month + 1, 1) - timedelta(days=1)
            bookings = bookings.filter(
                start_date__lte=last_day, end_date__gte=first_day
            ).exclude(booking_status=Booking.STATUS_CANCELLED)
        except (ValueError, TypeError):
            pass
    elif report_type == 'daily' and date_from_str:
        try:
            report_date = date.fromisoformat(date_from_str)
            bookings = bookings.filter(
                start_date__lte=report_date, end_date__gte=report_date
            ).exclude(booking_status=Booking.STATUS_CANCELLED)
        except ValueError:
            pass

    if date_from_str and report_type not in ('daily', 'monthly'):
        try:
            bookings = bookings.filter(start_date__gte=date.fromisoformat(date_from_str))
        except ValueError:
            pass

    if date_to_str and report_type not in ('daily', 'monthly'):
        try:
            bookings = bookings.filter(start_date__lte=date.fromisoformat(date_to_str))
        except ValueError:
            pass

    bookings = bookings.order_by('start_date', 'from_time')

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Bookings Report'

    # Header styling
    header_fill = PatternFill(start_color='4A3728', end_color='4A3728', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)
    header_alignment = Alignment(horizontal='center', vertical='center')

    headers = [
        'Booking #', 'Party Name', 'Contact', 'Function Type',
        'Start Date', 'End Date', 'Days', 'From Time', 'End Time',
        'Total (₹)', 'Received (₹)', 'Balance (₹)',
        'Payment Status', 'Booking Status', 'Remarks'
    ]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment

    # Data rows
    for row_idx, booking in enumerate(bookings, 2):
        ws.cell(row=row_idx, column=1, value=booking.booking_number)
        ws.cell(row=row_idx, column=2, value=booking.party_name)
        ws.cell(row=row_idx, column=3, value=booking.contact_number)
        ws.cell(row=row_idx, column=4, value=booking.get_function_type_display())
        ws.cell(row=row_idx, column=5, value=str(booking.start_date))
        ws.cell(row=row_idx, column=6, value=str(booking.end_date))
        ws.cell(row=row_idx, column=7, value=booking.number_of_days)
        ws.cell(row=row_idx, column=8, value=str(booking.from_time))
        ws.cell(row=row_idx, column=9, value=str(booking.end_time))
        ws.cell(row=row_idx, column=10, value=float(booking.total_payment))
        ws.cell(row=row_idx, column=11, value=float(booking.received_payment))
        ws.cell(row=row_idx, column=12, value=float(booking.balance))
        ws.cell(row=row_idx, column=13, value=booking.get_payment_status_display())
        ws.cell(row=row_idx, column=14, value=booking.get_booking_status_display())
        ws.cell(row=row_idx, column=15, value=booking.remarks or '')

    # Auto-fit columns
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 40)

    # Summary row
    summary_row = len(bookings) + 3
    ws.cell(row=summary_row, column=1, value='TOTALS').font = Font(bold=True)
    total_agg = bookings.aggregate(
        t=Sum('total_payment'),
        r=Sum('received_payment'),
        b=Sum('balance')
    )
    ws.cell(row=summary_row, column=10, value=float(total_agg['t'] or 0)).font = Font(bold=True)
    ws.cell(row=summary_row, column=11, value=float(total_agg['r'] or 0)).font = Font(bold=True)
    ws.cell(row=summary_row, column=12, value=float(total_agg['b'] or 0)).font = Font(bold=True)

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"bookings_report_{report_type}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Access-Control-Expose-Headers'] = 'Content-Disposition'
    return response


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_pdf(request):
    """Export bookings report as PDF."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    report_type = request.query_params.get('report_type', 'all')
    date_from_str = request.query_params.get('date_from')
    date_to_str = request.query_params.get('date_to')
    year_str = request.query_params.get('year')
    month_str = request.query_params.get('month')

    bookings = Booking.objects.all()

    if report_type == 'today':
        today = timezone.localdate()
        bookings = bookings.filter(start_date__lte=today, end_date__gte=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'upcoming':
        today = timezone.localdate()
        bookings = bookings.filter(start_date__gt=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'completed':
        today = timezone.localdate()
        bookings = bookings.filter(end_date__lt=today).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'cancelled':
        bookings = bookings.filter(booking_status=Booking.STATUS_CANCELLED)
    elif report_type == 'pending_payment':
        bookings = bookings.filter(
            payment_status__in=[Booking.PAYMENT_NOT_PAID, Booking.PAYMENT_PARTIAL]
        ).exclude(booking_status=Booking.STATUS_CANCELLED)
    elif report_type == 'fully_paid':
        bookings = bookings.filter(payment_status=Booking.PAYMENT_FULL).exclude(
            booking_status=Booking.STATUS_CANCELLED
        )
    elif report_type == 'monthly' and year_str and month_str:
        try:
            year = int(year_str)
            month = int(month_str)
            first_day = date(year, month, 1)
            if month == 12:
                last_day = date(year + 1, 1, 1) - timedelta(days=1)
            else:
                last_day = date(year, month + 1, 1) - timedelta(days=1)
            bookings = bookings.filter(
                start_date__lte=last_day, end_date__gte=first_day
            ).exclude(booking_status=Booking.STATUS_CANCELLED)
        except (ValueError, TypeError):
            pass

    bookings = bookings.order_by('start_date', 'from_time')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=0.5 * inch,
        leftMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle('Title', parent=styles['Heading1'], alignment=1, spaceAfter=12)
    story.append(Paragraph('Thiru Kumaran Mahal - Bookings Report', title_style))
    story.append(Paragraph(f'Report Type: {report_type.replace("_", " ").title()}', styles['Normal']))
    story.append(Paragraph(f'Generated: {timezone.now().strftime("%d/%m/%Y %H:%M")}', styles['Normal']))
    story.append(Spacer(1, 0.2 * inch))

    # Table headers
    header_color = colors.HexColor('#4A3728')
    data = [['Booking #', 'Party Name', 'Function', 'Start Date', 'End Date',
              'Days', 'Total', 'Received', 'Balance', 'Pay Status', 'Status']]

    for booking in bookings:
        data.append([
            booking.booking_number,
            booking.party_name[:20],
            booking.get_function_type_display()[:12],
            str(booking.start_date),
            str(booking.end_date),
            str(booking.number_of_days),
            f"Rs.{float(booking.total_payment):,.0f}",
            f"Rs.{float(booking.received_payment):,.0f}",
            f"Rs.{float(booking.balance):,.0f}",
            booking.get_payment_status_display()[:10],
            booking.get_booking_status_display()[:10],
        ])

    if len(data) > 1:
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), header_color),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F0EB')]),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
    else:
        story.append(Paragraph('No bookings found for this report.', styles['Normal']))

    doc.build(story)
    buffer.seek(0)

    filename = f"bookings_report_{report_type}_{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Access-Control-Expose-Headers'] = 'Content-Disposition'
    return response
